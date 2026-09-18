from pathlib import Path
import re
import sys


# ============================================================
# PROJECT PATH
# ============================================================

BASE_DIR = Path(__file__).resolve().parents[1]

if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))


# ============================================================
# DEPENDENCIES
# ============================================================

from openpyxl import load_workbook

from validation.validator import validate_workbook


# ============================================================
# PATHS
# ============================================================

WORKBOOK_PATH = BASE_DIR / "generated_tests_interview_ready_50.xlsx"

GENERATED_DIR = BASE_DIR / "tests" / "generated"

GENERATED_FILE = GENERATED_DIR / "test_generated_from_excel.py"

REPORT_DIR = BASE_DIR / "reports"

REPORT_FILE = REPORT_DIR / "excel_pytest_generation_report.txt"


# ============================================================
# SUPPORTED AUTOMATION MAPPINGS
# ============================================================

# These IDs correspond to scenarios that can currently be
# executed deterministically against the controller simulator.
#
# IMPORTANT:
# Do not invent simulator APIs for unsupported scenarios.
# Unsupported workbook rows are reported and skipped.

SUPPORTED_HANDLERS = {
    "TC_AI_003": "test_a2l_leak_shuts_down_compressor",
    "TC_AI_004": "test_a2l_leak_activates_ventilation_fan",
    "TC_AI_006": "test_a2l_leak_is_logged",
    "TC_AI_008": "test_manual_reset_clears_alarm_after_safe_condition",
    "TC_AI_017": "test_alarm_mode_activates_alarm",
    "TC_AI_018": "test_low_suction_pressure_shuts_down_compressor",
    "TC_AI_020": "test_invalid_sensor_input_enters_safe_mode",
    "TC_AI_027": "test_high_discharge_temperature_shuts_down_compressor",
}


MODBUS_HANDLERS = {
    "TC_AI_032": ("test_modbus_status_registers", 30004, 0),
    "TC_AI_035": ("test_modbus_status_registers", 30006, 0),
    "TC_AI_038": ("test_modbus_status_registers", 30007, 0),
}


MODBUS_LEAK_HANDLERS = {
    "TC_AI_039": "test_modbus_leak_status_register",
}


# ============================================================
# EXCEL READER
# ============================================================

def read_test_cases(workbook_path):
    """
    Read test cases from the Test Cases worksheet.

    Returns:
        list[dict]
    """

    workbook_path = Path(workbook_path)

    if not workbook_path.exists():
        raise FileNotFoundError(
            f"Workbook not found: {workbook_path}"
        )

    workbook = load_workbook(
        workbook_path,
        data_only=True,
    )

    if "Test Cases" not in workbook.sheetnames:
        raise ValueError(
            "Workbook does not contain a 'Test Cases' worksheet."
        )

    sheet = workbook["Test Cases"]

    headers = [
        cell.value
        for cell in sheet[1]
    ]

    if not headers:
        raise ValueError(
            "The 'Test Cases' worksheet has no headers."
        )

    rows = []

    for values in sheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        if not any(
            value is not None
            for value in values
        ):
            continue

        row = dict(
            zip(
                headers,
                values,
            )
        )

        rows.append(row)

    return rows


# ============================================================
# TEST NAME HELPERS
# ============================================================

def safe_test_name(test_case_id, scenario):
    """
    Convert a test-case ID and scenario into a safe pytest
    function name.
    """

    text = f"{test_case_id}_{scenario}"

    text = text.lower()

    text = re.sub(
        r"[^a-z0-9]+",
        "_",
        text,
    )

    text = text.strip("_")

    return f"test_{text}"


# ============================================================
# AUTOMATION ELIGIBILITY
# ============================================================

def is_automation_candidate(row):
    """
    Return True when the Excel row explicitly marks the test
    as an automation candidate.
    """

    value = row.get(
        "Automation Candidate",
        "",
    )

    return (
        str(value)
        .strip()
        .upper()
        == "YES"
    )


# ============================================================
# CONTROLLER TEST GENERATION
# ============================================================

def generate_controller_test(
    test_case_id,
    scenario=None,
):
    """
    Generate deterministic pytest code for a supported
    RefrigerationController scenario.

    The generated tests directly instantiate the simulator.

    They intentionally do NOT import existing pytest test
    functions from tests/test_controller_automation.py.
    """

    if test_case_id == "TC_AI_003":

        return '''
def test_tc_ai_003_a2l_leak_shuts_down_compressor():
    controller = RefrigerationController(
        compressor_on=True
    )

    controller.set_leak_condition()

    assert controller.compressor_on is False
'''.strip()

    if test_case_id == "TC_AI_004":

        return '''
def test_tc_ai_004_a2l_leak_activates_ventilation_fan():
    controller = RefrigerationController()

    controller.set_leak_condition()

    assert controller.evaporator_fan_on is True
'''.strip()

    if test_case_id == "TC_AI_006":

        return '''
def test_tc_ai_006_a2l_leak_is_logged():
    controller = RefrigerationController()

    controller.set_leak_condition()

    assert "refrigerant leak detected" in controller.event_log
'''.strip()

    if test_case_id == "TC_AI_008":

        return '''
def test_tc_ai_008_manual_reset_clears_alarm_after_safe_condition():
    controller = RefrigerationController()

    controller.set_leak_condition()

    assert controller.alarm_active is True

    controller.clear_leak_condition()
    controller.manual_reset()

    assert controller.alarm_active is False
'''.strip()

    if test_case_id == "TC_AI_017":

        return '''
def test_tc_ai_017_alarm_mode_activates_alarm():
    controller = RefrigerationController()

    controller.activate_alarm_mode()

    assert controller.alarm_mode is True
    assert controller.alarm_active is True
'''.strip()

    if test_case_id == "TC_AI_018":

        return '''
def test_tc_ai_018_low_suction_pressure_shuts_down_compressor():
    controller = RefrigerationController(
        compressor_on=True
    )

    controller.set_low_suction_pressure()

    assert controller.suction_pressure == "low"
    assert controller.compressor_on is False
'''.strip()

    if test_case_id == "TC_AI_020":

        return '''
def test_tc_ai_020_invalid_sensor_input_enters_safe_mode():
    controller = RefrigerationController()

    controller.invalidate_sensor()

    assert controller.sensor_input_valid is False
    assert controller.safe_mode is True
'''.strip()

    if test_case_id == "TC_AI_027":

        return '''
def test_tc_ai_027_high_discharge_temperature_shuts_down_compressor():
    controller = RefrigerationController(
        compressor_on=True
    )

    controller.set_high_discharge_temperature()

    assert controller.discharge_temperature == "above limit"
    assert controller.compressor_on is False
'''.strip()

    raise ValueError(
        f"Unsupported controller test case: {test_case_id}"
    )


# ============================================================
# MODBUS TEST GENERATION
# ============================================================

def generate_modbus_test(
    test_case_id,
    register,
    expected_value,
):
    """
    Generate a deterministic Modbus register test.
    """

    test_name = safe_test_name(
        test_case_id,
        f"modbus_register_{register}",
    )

    return f'''
def {test_name}():
    controller = RefrigerationController()

    registers = controller.modbus_registers()

    assert registers[{register}] == {expected_value!r}
'''.strip()


# ============================================================
# MODBUS LEAK TEST GENERATION
# ============================================================

def generate_modbus_leak_test(
    test_case_id,
):
    """
    Generate the Modbus A2L leak-status test.
    """

    return '''
def test_tc_ai_039_modbus_leak_status_register():
    controller = RefrigerationController()

    controller.set_leak_condition()

    registers = controller.modbus_registers()

    assert registers[30007] == 1
'''.strip()


# ============================================================
# GENERATED TEST FILE
# ============================================================

def build_generated_file(test_blocks):
    """
    Build the complete generated pytest module.
    """

    header = '''"""
AUTO-GENERATED FILE.

Generated from:
    generated_tests_interview_ready_50.xlsx

Do not manually edit this file.

The source of truth is the validated Excel workbook.
"""

from simulators.controller_simulator import (
    RefrigerationController,
)


'''

    return header + "\n\n\n".join(
        test_blocks
    ) + "\n"


# ============================================================
# GENERATION
# ============================================================

def generate_pytest_from_excel(
    workbook_path=WORKBOOK_PATH,
):
    """
    Validate the workbook, identify automation candidates,
    generate supported deterministic pytest tests, and
    report unsupported scenarios.
    """

    workbook_path = Path(workbook_path)

    print(
        f"Validating workbook: {workbook_path.name}"
    )

    # --------------------------------------------------------
    # STEP 1 - VALIDATE WORKBOOK
    # --------------------------------------------------------

    validate_workbook(
        str(workbook_path)
    )

    print("Workbook validation passed.")

    # --------------------------------------------------------
    # STEP 2 - READ TEST CASES
    # --------------------------------------------------------

    rows = read_test_cases(
        workbook_path
    )

    # --------------------------------------------------------
    # STEP 3 - GENERATE TESTS
    # --------------------------------------------------------

    generated_blocks = []

    generated_cases = []
    skipped_cases = []

    seen_ids = set()

    for row in rows:

        test_case_id = str(
            row.get(
                "Test Case ID",
                "",
            )
        ).strip()

        scenario = str(
            row.get(
                "Scenario",
                "",
            )
        ).strip()

        if not test_case_id:
            continue

        # Prevent duplicate Excel IDs from producing duplicate
        # pytest functions.
        if test_case_id in seen_ids:
            skipped_cases.append(
                (
                    test_case_id,
                    scenario,
                    "Duplicate test-case ID",
                )
            )
            continue

        seen_ids.add(test_case_id)

        # ----------------------------------------------------
        # AUTOMATION CANDIDATE CHECK
        # ----------------------------------------------------

        if not is_automation_candidate(row):

            skipped_cases.append(
                (
                    test_case_id,
                    scenario,
                    "Automation Candidate is not YES",
                )
            )

            continue

        # ----------------------------------------------------
        # STANDARD CONTROLLER TEST
        # ----------------------------------------------------

        if test_case_id in SUPPORTED_HANDLERS:

            try:
                code = generate_controller_test(
                    test_case_id,
                    scenario,
                )

                generated_blocks.append(code)

                generated_cases.append(
                    (
                        test_case_id,
                        scenario,
                    )
                )

            except ValueError as exc:

                skipped_cases.append(
                    (
                        test_case_id,
                        scenario,
                        str(exc),
                    )
                )

            continue

        # ----------------------------------------------------
        # MODBUS REGISTER TEST
        # ----------------------------------------------------

        if test_case_id in MODBUS_HANDLERS:

            (
                _handler_name,
                register,
                expected_value,
            ) = MODBUS_HANDLERS[
                test_case_id
            ]

            code = generate_modbus_test(
                test_case_id,
                register,
                expected_value,
            )

            generated_blocks.append(code)

            generated_cases.append(
                (
                    test_case_id,
                    scenario,
                )
            )

            continue

        # ----------------------------------------------------
        # MODBUS LEAK STATUS TEST
        # ----------------------------------------------------

        if test_case_id in MODBUS_LEAK_HANDLERS:

            code = generate_modbus_leak_test(
                test_case_id,
            )

            generated_blocks.append(code)

            generated_cases.append(
                (
                    test_case_id,
                    scenario,
                )
            )

            continue

        # ----------------------------------------------------
        # UNSUPPORTED
        # ----------------------------------------------------

        skipped_cases.append(
            (
                test_case_id,
                scenario,
                "Scenario is not currently supported by the deterministic simulator",
            )
        )

    # --------------------------------------------------------
    # STEP 4 - WRITE GENERATED PYTEST MODULE
    # --------------------------------------------------------

    GENERATED_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    generated_content = build_generated_file(
        generated_blocks
    )

    GENERATED_FILE.write_text(
        generated_content,
        encoding="utf-8",
    )

    # --------------------------------------------------------
    # STEP 5 - WRITE REPORT
    # --------------------------------------------------------

    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    report_lines = [
        "Excel to Pytest Generation Report",
        "=" * 60,
        "",
        f"Workbook: {workbook_path}",
        f"Generated file: {GENERATED_FILE}",
        "",
        f"Total workbook cases: {len(rows)}",
        f"Generated tests: {len(generated_cases)}",
        f"Skipped tests: {len(skipped_cases)}",
        "",
        "Generated Test Cases",
        "-" * 60,
    ]

    if generated_cases:

        for test_case_id, scenario in generated_cases:

            report_lines.append(
                f"{test_case_id} | {scenario}"
            )

    else:

        report_lines.append(
            "None"
        )

    report_lines.extend(
        [
            "",
            "Skipped Test Cases",
            "-" * 60,
        ]
    )

    if skipped_cases:

        for (
            test_case_id,
            scenario,
            reason,
        ) in skipped_cases:

            report_lines.append(
                f"{test_case_id} | {scenario} | {reason}"
            )

    else:

        report_lines.append(
            "None"
        )

    REPORT_FILE.write_text(
        "\n".join(report_lines) + "\n",
        encoding="utf-8",
    )

    # --------------------------------------------------------
    # STEP 6 - CONSOLE SUMMARY
    # --------------------------------------------------------

    print(
        f"Generated: {GENERATED_FILE}"
    )

    print(
        f"Report:    {REPORT_FILE}"
    )

    print(
        f"Generated tests: {len(generated_cases)}"
    )

    print(
        f"Skipped tests:   {len(skipped_cases)}"
    )

    return {
        "total": len(rows),
        "generated": len(generated_cases),
        "skipped": len(skipped_cases),
        "generated_file": GENERATED_FILE,
        "report_file": REPORT_FILE,
        "generated_cases": generated_cases,
        "skipped_cases": skipped_cases,
    }


# ============================================================
# COMMAND-LINE ENTRY POINT
# ============================================================

if __name__ == "__main__":

    generate_pytest_from_excel()