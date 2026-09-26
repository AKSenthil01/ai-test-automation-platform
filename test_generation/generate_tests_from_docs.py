
"""
Generate controlled pytest adapters from the reviewed Excel test workbook.

Phase 1 of Excel -> pytest generation:
- Validate the workbook first.
- Read the Automation sheet.
- Match each scenario to an approved simulator capability.
- Generate executable pytest code from the capability, not from the test ID.
- Report supported and unsupported cases.
- Never invent simulator APIs.

Capability matching is deterministic and uses:
    Feature
    Scenario
    Steps
    Expected Result
    Test Data

Test Case IDs are used only for generated pytest function names.

Run from the project root:

    python test_generation/generate_tests_from_docs.py

or:

    python test_generation/generate_tests_from_docs.py generated_tests.xlsx

Output:

    tests/generated/test_generated_from_excel.py
    reports/excel_pytest_generation_report.txt
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl


# ---------------------------------------------------------------------------
# Project path
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent.parent

if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))


from validation.validator import validate_workbook


# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

DEFAULT_WORKBOOK = BASE_DIR / "generated_tests.xlsx"

OUTPUT_DIR = BASE_DIR / "tests" / "generated"
OUTPUT_FILE = OUTPUT_DIR / "test_generated_from_excel.py"

REPORT_DIR = BASE_DIR / "reports"
REPORT_FILE = REPORT_DIR / "excel_pytest_generation_report.txt"


# ---------------------------------------------------------------------------
# Capability model
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Capability:
    """Approved deterministic simulator capability."""

    name: str
    description: str
    action: str
    register: int | None = None
    expected: object | None = None


# ---------------------------------------------------------------------------
# Approved simulator capabilities
# ---------------------------------------------------------------------------

CAPABILITIES = {
    "a2l_compressor_shutdown": Capability(
        name="a2l_compressor_shutdown",
        description="A2L leak condition shuts down the compressor",
        action="set_leak_condition",
    ),

    "a2l_ventilation": Capability(
        name="a2l_ventilation",
        description="A2L leak condition activates ventilation",
        action="set_leak_condition",
    ),

    "a2l_alarm": Capability(
        name="a2l_alarm",
        description="A2L leak condition activates the alarm",
        action="set_leak_condition",
    ),

    "a2l_logging": Capability(
        name="a2l_logging",
        description="A2L leak condition is logged",
        action="set_leak_condition",
    ),

    "manual_reset": Capability(
        name="manual_reset",
        description="Manual reset clears the alarm after safe condition",
        action="set_leak_condition",
    ),

    "alarm_mode": Capability(
        name="alarm_mode",
        description="Controller enters alarm mode",
        action="activate_alarm_mode",
    ),

    "low_suction_pressure": Capability(
        name="low_suction_pressure",
        description="Low suction pressure protection shuts down compressor",
        action="set_low_suction_pressure",
    ),

    "invalid_sensor_safe_mode": Capability(
        name="invalid_sensor_safe_mode",
        description="Invalid sensor input causes safe mode",
        action="invalidate_sensor",
    ),

    "high_discharge_temperature": Capability(
        name="high_discharge_temperature",
        description="High discharge temperature protection shuts down compressor",
        action="set_high_discharge_temperature",
    ),
}


# ---------------------------------------------------------------------------
# Modbus register semantics
# ---------------------------------------------------------------------------

MODBUS_REGISTER_MAP = {
    30001: (
        "evaporator temperature",
        "evaporator temperature",
    ),
    30002: (
        "suction pressure",
        "suction pressure",
    ),
    30003: (
        "discharge temperature",
        "discharge temperature",
    ),
    30004: (
        "compressor status",
        0,
    ),
    30006: (
        "alarm status",
        0,
    ),
    30007: (
        "a2l leak status",
        0,
    ),
    30008: (
        "fan status",
        0,
    ),
}


# ---------------------------------------------------------------------------
# Workbook loading
# ---------------------------------------------------------------------------


def load_automation_rows(workbook_path: Path) -> list[dict[str, str]]:
    """Load automation candidates from either supported workbook schema."""

    workbook = openpyxl.load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    try:
        preferred_sheets = []

        if "Automation" in workbook.sheetnames:
            preferred_sheets.append("Automation")

        if "Test Cases" in workbook.sheetnames:
            preferred_sheets.append("Test Cases")

        if not preferred_sheets:
            raise ValueError(
                "Workbook must contain either 'Automation' "
                "or 'Test Cases' sheet."
            )

        required_columns = {
            "Test Case ID",
            "Feature",
            "Scenario",
            "Steps",
            "Test Data",
            "Automation Candidate",
            "Tool",
        }

        expected_result_columns = {
            "Expected Result",
            "Expected Results",
        }

        selected_sheet = None
        headers = None

        for sheet_name in preferred_sheets:

            worksheet = workbook[sheet_name]

            rows = worksheet.iter_rows(
                values_only=True
            )

            try:
                header_row = next(rows)
            except StopIteration:
                continue

            normalized_headers = [
                str(value).strip()
                if value is not None
                else ""
                for value in header_row
            ]

            header_set = set(normalized_headers)

            missing_required = (
                required_columns - header_set
            )

            has_expected_result = bool(
                expected_result_columns & header_set
            )

            if not missing_required and has_expected_result:
                selected_sheet = worksheet
                headers = normalized_headers
                break

        if selected_sheet is None:
            available = []

            for sheet_name in preferred_sheets:
                worksheet = workbook[sheet_name]

                rows = worksheet.iter_rows(
                    values_only=True
                )

                try:
                    header_row = next(rows)
                except StopIteration:
                    available.append(
                        f"{sheet_name}: empty"
                    )
                    continue

                normalized_headers = [
                    str(value).strip()
                    if value is not None
                    else ""
                    for value in header_row
                ]

                available.append(
                    f"{sheet_name}: {normalized_headers}"
                )

            raise ValueError(
                "No workbook sheet contains the required "
                "automation columns plus 'Expected Result' "
                "or 'Expected Results'.\n"
                + "\n".join(available)
            )

        print(
            f"Reading workbook sheet: "
            f"{selected_sheet.title}"
        )

        header_index = {
            header: index
            for index, header in enumerate(headers)
        }

        expected_result_header = next(
            header
            for header in expected_result_columns
            if header in header_index
        )

        records = []

        for row in selected_sheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            if not any(
                value is not None
                for value in row
            ):
                continue

            record = {}

            for column in required_columns:
                index = header_index[column]

                value = (
                    row[index]
                    if index < len(row)
                    else ""
                )

                record[column] = (
                    str(value).strip()
                    if value is not None
                    else ""
                )

            expected_index = header_index[
                expected_result_header
            ]

            expected_value = (
                row[expected_index]
                if expected_index < len(row)
                else ""
            )

            record["Expected Result"] = (
                str(expected_value).strip()
                if expected_value is not None
                else ""
            )

            records.append(record)

        return records

    finally:
        workbook.close()

# ---------------------------------------------------------------------------
# Text normalization
# ---------------------------------------------------------------------------

def normalize(text: str) -> str:
    """Basic lowercase and whitespace normalization."""

    text = str(text or "").lower()

    text = text.replace("_", " ")
    text = text.replace("-", " ")

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip()


def normalize_for_matching(text: str) -> str:
    """
    Normalize common HVAC wording variations.

    Examples:

        shuts down -> shutdown
        shut down -> shutdown
        shuts off -> off
        is off -> off
        fan activates -> fan activate
        fans activate -> fan activate
        leak detection -> leak
        sensor failure -> invalid sensor
    """

    text = normalize(text)

    # ---------------------------------------------------------------
    # Compressor shutdown wording
    # ---------------------------------------------------------------

    text = re.sub(
        r"\bshuts?\s+down\b",
        "shutdown",
        text,
    )

    text = re.sub(
        r"\bshut\s+down\b",
        "shutdown",
        text,
    )

    text = re.sub(
        r"\bshuts?\s+off\b",
        "off",
        text,
    )

    text = re.sub(
        r"\bturns?\s+off\b",
        "off",
        text,
    )

    text = re.sub(
        r"\bswitch(?:es)?\s+off\b",
        "off",
        text,
    )

    text = re.sub(
        r"\bis\s+off\b",
        "off",
        text,
    )

    text = re.sub(
        r"\bstate\s+is\s+off\b",
        "off",
        text,
    )

    # ---------------------------------------------------------------
    # Fan / ventilation wording
    # ---------------------------------------------------------------

    text = re.sub(
        r"\bfans?\s+activate\b",
        "fan activate",
        text,
    )

    text = re.sub(
        r"\bactivate(?:s)?\s+fans?\b",
        "fan activate",
        text,
    )

    text = re.sub(
        r"\bfans?\s+activated\b",
        "fan activate",
        text,
    )

    text = re.sub(
        r"\bventilat(?:e|es|ed|ing|ion)\b",
        "ventilation",
        text,
    )

    # ---------------------------------------------------------------
    # Leak wording
    # ---------------------------------------------------------------

    text = re.sub(
        r"\bleak\s+detection\b",
        "leak",
        text,
    )

    text = re.sub(
        r"\bleak\s+detected\b",
        "leak",
        text,
    )

    text = re.sub(
        r"\bdetected\s+leak\b",
        "leak",
        text,
    )

    # ---------------------------------------------------------------
    # Sensor wording
    # ---------------------------------------------------------------

    text = re.sub(
        r"\bsensor\s+failure\b",
        "invalid sensor",
        text,
    )

    text = re.sub(
        r"\bsensor\s+input\s+invalid\b",
        "invalid sensor",
        text,
    )

    # ---------------------------------------------------------------
    # Temperature / pressure wording
    # ---------------------------------------------------------------

    text = re.sub(
        r"\bhigh\s+discharge\s+temp\b",
        "high discharge temperature",
        text,
    )

    text = re.sub(
        r"\blow\s+suction\b",
        "low suction",
        text,
    )

    return text


# ---------------------------------------------------------------------------
# Record helpers
# ---------------------------------------------------------------------------

def safe_name(test_case_id: str) -> str:
    """Convert Excel test-case ID to a safe pytest function suffix."""

    return re.sub(
        r"[^a-zA-Z0-9_]",
        "_",
        test_case_id.lower(),
    )


def combined_record_text(
    record: dict[str, str],
) -> str:
    """
    Combine all semantic Excel fields.

    Test Case ID is deliberately excluded.
    """

    fields = [
        "Feature",
        "Scenario",
        "Steps",
        "Expected Result",
        "Test Data",
    ]

    return normalize_for_matching(
        " ".join(
            record.get(
                field,
                "",
            )
            for field in fields
        )
    )


def combined_scenario_text(
    record: dict[str, str],
) -> str:
    """
    Backward-compatible helper.

    Capability matching now uses the complete Excel record.
    """

    return combined_record_text(record)


def _contains_all(
    text: str,
    *terms: str,
) -> bool:
    return all(
        normalize_for_matching(term) in text
        for term in terms
    )


def _contains_any(
    text: str,
    *terms: str,
) -> bool:
    return any(
        normalize_for_matching(term) in text
        for term in terms
    )


# ---------------------------------------------------------------------------
# Register helpers
# ---------------------------------------------------------------------------

def _extract_register(
    text: str,
) -> int | None:
    """Extract supported Modbus register numbers."""

    matches = re.findall(
        r"\b3000[1-8]\b",
        text,
    )

    if not matches:
        return None

    return int(matches[0])


def _extract_status_value(
    text: str,
) -> int | None:
    """
    Extract explicit Modbus 0/1 values.

    Supports:

        value 1
        status 1
        value is 1
        status = 0
        register 30007 to 1
    """

    text = normalize(text)

    patterns = [
        r"\b(?:status|value)\s*(?:value\s*)?(?:is|=|to)?\s*([01])\b",
        r"\b(?:value|status)\s+([01])\b",
        r"\bto\s+([01])\b",
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
        )

        if match:
            return int(
                match.group(1)
            )

    return None


# ---------------------------------------------------------------------------
# Semantic capability helpers
# ---------------------------------------------------------------------------

def _has_compressor_shutdown(
    text: str,
) -> bool:
    """
    Detect compressor shutdown semantics.

    Handles:

        compressor shutdown
        compressor shuts down
        compressor shut down
        compressor turns off
        compressor is off
        compressor state is off
    """

    if "compressor" not in text:
        return False

    return _contains_any(
        text,
        "compressor shutdown",
        "compressor off",
        "compressor state off",
        "compressor turns off",
        "compressor shut down",
    )


def _has_fan_activation(
    text: str,
) -> bool:
    """Detect ventilation/fan activation semantics."""

    return (
        _contains_any(
            text,
            "fan activate",
            "fan activation",
            "fan on",
            "fan state on",
            "fans on",
        )
        or (
            "ventilation" in text
            and "fan" in text
        )
    )


def _has_a2l_leak(
    text: str,
) -> bool:
    """Detect A2L refrigerant leak semantics."""

    return (
        "a2l" in text
        and "leak" in text
    )


def _has_event_logging(
    text: str,
) -> bool:
    """Detect event logging semantics."""

    return _contains_any(
        text,
        "event logged",
        "event logging",
        "logged in controller memory",
        "event log",
        "logging",
    )


def _has_manual_reset(
    text: str,
) -> bool:
    """
    Detect manual reset after a safe condition.

    The simulator requires:
        leak condition
        clear leak condition
        manual reset
    """

    return (
        _contains_all(
            text,
            "manual",
            "reset",
        )
        and _contains_any(
            text,
            "safe condition",
            "below safe limit",
            "concentration drops below",
            "normal operation",
        )
    )


def _has_alarm_mode(
    text: str,
) -> bool:
    """Detect explicit alarm-mode activation."""

    return (
        "alarm mode" in text
        and _contains_any(
            text,
            "activate",
            "activated",
            "enter",
            "enters",
            "enable",
        )
    )


def _has_low_suction_pressure(
    text: str,
) -> bool:
    """Detect low suction pressure protection."""

    return (
        _contains_all(
            text,
            "low",
            "suction",
            "pressure",
        )
        and _has_compressor_shutdown(text)
    )


def _has_invalid_sensor_safe_mode(
    text: str,
) -> bool:
    """Detect invalid sensor -> safe mode."""

    return (
        _contains_all(
            text,
            "invalid",
            "sensor",
        )
        and "safe mode" in text
    )


def _has_high_discharge_temperature(
    text: str,
) -> bool:
    """Detect high discharge temperature protection."""

    return (
        _contains_all(
            text,
            "high",
            "discharge",
            "temperature",
        )
        and _has_compressor_shutdown(text)
    )


# ---------------------------------------------------------------------------
# Capability matcher
# ---------------------------------------------------------------------------

def match_capability(
    record: dict[str, str],
) -> Capability | None:
    """
    Match an Excel test case to an approved simulator capability.

    Test Case ID is never used to select the simulator behavior.

    Matching is based on:
        Feature
        Scenario
        Expected Result
        Steps
        Test Data

    The matcher deliberately avoids using Requirement as the primary
    capability selector because a requirement may describe several
    behaviors while an individual test case normally verifies one
    behavior.
    """

    feature = normalize_for_matching(
        record.get("Feature", "")
    )

    scenario = normalize_for_matching(
        record.get("Scenario", "")
    )

    expected = normalize_for_matching(
        record.get("Expected Result", "")
    )

    steps = normalize_for_matching(
        record.get("Steps", "")
    )

    test_data = normalize_for_matching(
        record.get("Test Data", "")
    )

    primary = " ".join(
        [
            feature,
            scenario,
            expected,
        ]
    )

    supporting = " ".join(
        [
            steps,
            test_data,
        ]
    )

    full_text = " ".join(
        [
            primary,
            supporting,
        ]
    )

    # ===============================================================
    # 1. Explicit unsupported / negative alarm-mode scenarios
    # ===============================================================

    alarm_not_active = _contains_any(
        primary,
        "alarm mode is not activated",
        "alarm mode is not active",
        "alarm mode remains inactive",
        "alarm should not activate",
        "alarm is not activated",
        "alarm is not active",
        "no alarm",
    )

    if alarm_not_active:
        return None

    # ===============================================================
    # 2. Complex alarm-mode scenarios
    # ===============================================================
    #
    # Do not map a complex scenario to only one partial behavior.
    #
    # Example:
    #   high discharge temperature
    #   + alarm mode activation
    #   + compressor shutdown
    #
    # The simulator does not implement that complete combined
    # sequence, so do not pretend that it does.
    # ===============================================================

    complex_alarm_mode = (
        "alarm mode" in primary
        and _contains_any(
            primary,
            "high discharge temperature",
            "compressor shutdown",
            "compressor shut down",
            "abnormal condition",
        )
    )

    if complex_alarm_mode:
        return None

    # ===============================================================
    # 3. Modbus register tests
    # ===============================================================
    #
    # Check Modbus before A2L because a Modbus leak-status test can
    # contain A2L + leak terminology.
    # ===============================================================

    register = _extract_register(
        full_text
    )

    if (
        "modbus" in full_text
        or "register" in full_text
        or register is not None
    ):

        if register == 30005:
            # Defrost register intentionally unsupported.
            return None

        if register in MODBUS_REGISTER_MAP:

            label, default_expected = (
                MODBUS_REGISTER_MAP[register]
            )

            expected_value = _extract_status_value(
                full_text
            )

            if expected_value is None:
                expected_value = default_expected

            return Capability(
                name=f"modbus_register_{register}",
                description=(
                    f"Modbus register {register} -> {label}"
                ),
                action="modbus_registers",
                register=register,
                expected=expected_value,
            )

    # ===============================================================
    # 4. A2L context
    # ===============================================================

    a2l = (
        _contains_any(
            feature,
            "a2l",
            "leak detection",
        )
        or _contains_any(
            scenario,
            "a2l",
            "refrigerant leak",
            "leak detection",
            "leak detected",
            "refrigerant concentration",
        )
    )

    if a2l:

        # -----------------------------------------------------------
        # A2L -> compressor shutdown
        # -----------------------------------------------------------

        if _contains_any(
            primary,
            "compressor shutdown",
            "compressor shuts down",
            "compressor shut down",
            "immediate compressor shutdown",
            "compressor is shut down",
            "compressor should be off",
            "compressor must be off",
            "compressor turns off",
        ):
            return CAPABILITIES[
                "a2l_compressor_shutdown"
            ]

        # -----------------------------------------------------------
        # A2L -> ventilation
        # -----------------------------------------------------------

        if _contains_any(
            primary,
            "ventilation",
            "ventilation fan",
            "fan activates to ventilate",
            "fan activation",
            "evaporator fan activates",
            "fans activate to ventilate",
            "fan is activated",
        ):
            return CAPABILITIES[
                "a2l_ventilation"
            ]

        # -----------------------------------------------------------
        # A2L -> alarm
        # -----------------------------------------------------------

        if _contains_any(
            primary,
            "alarm activates",
            "alarm activated",
            "alarm status activated",
            "alarm status activation",
            "alarm becomes active",
            "alarm is active",
            "alarm should be active",
            "leak alarm activates",
            "leak alarm",
        ):
            return CAPABILITIES[
                "a2l_alarm"
            ]

        # -----------------------------------------------------------
        # A2L -> event logging
        # -----------------------------------------------------------

        if _contains_any(
            primary,
            "event logged",
            "event is logged",
            "event logging",
            "leak is logged",
            "leak event logging",
            "logged in controller memory",
            "event log",
        ):
            return CAPABILITIES[
                "a2l_logging"
            ]

    # ===============================================================
    # 5. Manual reset
    # ===============================================================

    if _contains_any(
        primary,
        "manual reset",
        "manual reset clears",
        "manual reset performed",
        "reset performed by technician",
        "technician performs manual reset",
        "manual reset after safe condition",
    ):
        return CAPABILITIES[
            "manual_reset"
        ]

    if (
        "manual reset" in scenario
        or "reset after safe condition" in scenario
    ):
        return CAPABILITIES[
            "manual_reset"
        ]

    # ===============================================================
    # 6. High discharge temperature
    # ===============================================================

    if _contains_any(
        primary,
        "high discharge temperature",
        "discharge temperature protection",
        "discharge temperature exceeds limit",
        "discharge temperature above limit",
        "high discharge protection",
    ):
        return CAPABILITIES[
            "high_discharge_temperature"
        ]

    # ===============================================================
    # 7. Low suction pressure
    # ===============================================================

    if _contains_any(
        primary,
        "low suction pressure",
        "suction pressure protection",
        "suction pressure is low",
        "low suction protection",
    ):
        return CAPABILITIES[
            "low_suction_pressure"
        ]

    # ===============================================================
    # 8. Invalid sensor -> safe mode
    # ===============================================================

    if _contains_any(
        primary,
        "invalid sensor input",
        "sensor input is invalid",
        "sensor failure protection",
        "invalid sensor",
        "enters safe mode",
        "enter safe mode",
        "controller enters safe mode",
        "safe mode after invalid sensor",
    ):
        return CAPABILITIES[
            "invalid_sensor_safe_mode"
        ]

    # ===============================================================
    # 9. Simple alarm mode
    # ===============================================================

    if _contains_any(
        scenario,
        "alarm mode activation",
        "activate alarm mode",
        "controller enters alarm mode",
    ):
        return CAPABILITIES[
            "alarm_mode"
        ]

    # ===============================================================
    # 10. No approved simulator capability
    # ===============================================================

    return None
# ---------------------------------------------------------------------------
# Generated pytest code
# ---------------------------------------------------------------------------

def build_test_code(
    supported: list[
        tuple[
            dict[str, str],
            Capability,
        ]
    ],
) -> str:
    """Generate pytest source from capability mappings."""

    lines = [
        '"""Generated pytest tests from the validated Excel workbook.',
        "",
        "DO NOT EDIT MANUALLY.",
        "Regenerate with:",
        "python test_generation/generate_tests_from_docs.py generated_tests.xlsx",
        '"""',
        "",
        "from simulators.controller_simulator import "
        "RefrigerationController",
        "",
    ]

    for record, capability in supported:

        test_case_id = record[
            "Test Case ID"
        ]

        generated_name = (
            f"test_generated_{safe_name(test_case_id)}"
        )

        lines.extend(
            [
                f"def {generated_name}():",
                (
                    f'    """Generated from Excel '
                    f'test case {test_case_id}."""'
                ),
                "    controller = RefrigerationController()",
            ]
        )

        # -----------------------------------------------------------
        # Modbus register capability
        # -----------------------------------------------------------

        if capability.action == "modbus_registers":

            setup_actions = {
                (30004, 1):
                    "    controller.compressor_on = True",

                (30006, 1):
                    "    controller.activate_alarm_mode()",

                (30007, 1):
                    "    controller.set_leak_condition()",

                (30008, 1):
                    "    controller.evaporator_fan_on = True",
            }

            setup = setup_actions.get(
                (
                    capability.register,
                    capability.expected,
                )
            )

            if setup:
                lines.append(setup)

            lines.append(
                "    registers = "
                "controller.modbus_registers()"
            )

            lines.append(
                f"    assert registers[{capability.register}] "
                f"== {capability.expected!r}"
            )

        # -----------------------------------------------------------
        # Normal simulator capability
        # -----------------------------------------------------------

        else:

            lines.append(
                f"    controller.{capability.action}()"
            )

            assertions = {
                "a2l_compressor_shutdown":
                    "    assert "
                    "controller.compressor_on is False",

                "a2l_ventilation":
                    "    assert "
                    "controller.evaporator_fan_on is True",

                "a2l_alarm": " assert " "controller.alarm_active is True",

                "a2l_logging":
                    "    assert "
                    "'refrigerant leak detected' "
                    "in controller.event_log",

                "manual_reset":
                    (
                        "    controller.clear_leak_condition()\n"
                        "    controller.manual_reset()\n"
                        "    assert controller.alarm_active is False"
                    ),

                "alarm_mode":
                    "    assert "
                    "controller.alarm_active is True",

                "low_suction_pressure":
                    "    assert "
                    "controller.compressor_on is False",

                "invalid_sensor_safe_mode":
                    "    assert "
                    "controller.safe_mode is True",

                "high_discharge_temperature":
                    "    assert "
                    "controller.compressor_on is False",
            }

            assertion = assertions.get(
                capability.name
            )

            if assertion is None:
                raise RuntimeError(
                    "No pytest assertion is defined for "
                    f"capability: {capability.name}"
                )


            # Ensure every generated assertion line is indented
            # exactly four spaces inside the pytest function.
            for assertion_line in assertion.splitlines():
                lines.append(
                    "    " + assertion_line.strip()
                )


        lines.append("")

    if not supported:

        lines.extend(
            [
                "def test_no_supported_generated_cases():",
                (
                    '    pytest.skip('
                    '"No supported simulator capabilities were found."'
                    ")"
                ),
                "",
            ]
        )

    return (
        "\n".join(lines)
        .rstrip()
        + "\n"
    )


# ---------------------------------------------------------------------------
# Generation pipeline
# ---------------------------------------------------------------------------

def generate(
    workbook: Path,
) -> tuple[
    list[str],
    list[tuple[str, str]],
]:

    workbook = workbook.resolve()

    if not workbook.exists():
        raise FileNotFoundError(
            f"Workbook not found: {workbook}"
        )

    print(
        f"Validating workbook: {workbook.name}"
    )

    validate_workbook(
        str(workbook)
    )

    print(
        "Workbook validation passed."
    )

    records = load_automation_rows(
        workbook
    )

    supported: list[
        tuple[
            dict[str, str],
            Capability,
        ]
    ] = []

    skipped: list[
        tuple[
            str,
            str,
        ]
    ] = []

    # ---------------------------------------------------------------
    # Capability matching
    # ---------------------------------------------------------------

    print()
    print(
        "Capability matching:"
    )

    for record in records:

        test_case_id = record[
            "Test Case ID"
        ]

        candidate = normalize(
            record.get(
                "Automation Candidate",
                "",
            )
        )

        tool = normalize(
            record.get(
                "Tool",
                "",
            )
        )

        if candidate not in {
            "yes",
            "y",
            "true",
        }:

            skipped.append(
                (
                    test_case_id,
                    "Automation Candidate is not YES",
                )
            )

            print(
                f"  SKIP {test_case_id} -> "
                "Automation Candidate is not YES"
            )

            continue

        if tool != "pytest":

            skipped.append(
                (
                    test_case_id,
                    f"Unsupported tool: "
                    f"{record.get('Tool')}",
                )
            )

            print(
                f"  SKIP {test_case_id} -> "
                f"unsupported tool: "
                f"{record.get('Tool')}"
            )

            continue

        capability = match_capability(
            record
        )

        if capability is None:

            skipped.append(
                (
                    test_case_id,
                    (
                        "No approved simulator capability "
                        "matches the scenario"
                    ),
                )
            )

            print(
                f"  SKIP {test_case_id} -> "
                "no approved simulator capability"
            )

            continue

        supported.append(
            (
                record,
                capability,
            )
        )

        print(
            f"  MAP  {test_case_id} -> "
            f"{capability.name} "
            f"({capability.action})"
        )

    # ---------------------------------------------------------------
    # Output directories
    # ---------------------------------------------------------------

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    REPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # ---------------------------------------------------------------
    # Generated pytest
    # ---------------------------------------------------------------

    code = build_test_code(
        supported
    )

    OUTPUT_FILE.write_text(
        code,
        encoding="utf-8",
    )

    # ---------------------------------------------------------------
    # Report
    # ---------------------------------------------------------------

    report_lines = [
        "Excel -> Pytest Generation Report",
        "=" * 40,
        f"Workbook: {workbook.name}",
        f"Automation rows: {len(records)}",
        f"Generated pytest tests: {len(supported)}",
        f"Skipped/unsupported: {len(skipped)}",
        "",
        "Generated capability mappings:",
        "-" * 40,
    ]

    for record, capability in supported:

        report_lines.append(
            f"  {record['Test Case ID']} | "
            f"{capability.name} | "
            f"{capability.action} | "
            f"{record['Scenario']}"
        )

    report_lines.append("")
    report_lines.append(
        "Skipped/unsupported:"
    )
    report_lines.append(
        "-" * 40
    )

    if skipped:

        for test_case_id, reason in skipped:

            report_lines.append(
                f"  {test_case_id} | {reason}"
            )

    else:

        report_lines.append(
            "  None"
        )

    REPORT_FILE.write_text(
        "\n".join(report_lines)
        + "\n",
        encoding="utf-8",
    )

    # ---------------------------------------------------------------
    # Console summary
    # ---------------------------------------------------------------

    print()
    print(
        f"Generated: {OUTPUT_FILE}"
    )

    print(
        f"Report:    {REPORT_FILE}"
    )

    print(
        f"Automation rows: {len(records)}"
    )

    print(
        f"Generated tests: {len(supported)}"
    )

    print(
        f"Skipped tests:   {len(skipped)}"
    )

    return (
        [
            record["Test Case ID"]
            for record, _ in supported
        ],
        skipped,
    )


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> int:

    parser = argparse.ArgumentParser(
        description=(
            "Generate controlled pytest tests "
            "from a validated Excel workbook."
        )
    )

    parser.add_argument(
        "workbook",
        nargs="?",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help=(
            "Path to the Excel workbook. "
            "Defaults to generated_tests.xlsx."
        ),
    )

    args = parser.parse_args()

    workbook = args.workbook

    if not workbook.is_absolute():
        workbook = BASE_DIR / workbook

    try:

        generate(
            workbook
        )

    except Exception as exc:

        print(
            f"ERROR: {exc}"
        )

        return 1

    return 0


if __name__ == "__main__":
    raise SystemExit(
        main()
    )
